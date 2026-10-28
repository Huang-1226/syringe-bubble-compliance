# -*- coding: utf-8 -*-
"""Build the English technical report on bubble compliance."""

import os
import re
from copy import deepcopy

import pandas as pd
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(BASE, "analysis", "data", "bubble")
OUT = os.path.join(BASE, "report", "Bubble_Compliance_Report_EN.docx")

# ---------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------

SUB_TOKENS = [
    ("K_eff", "K", "eff"),
    ("K_air", "K", "air"),
    ("A_out", "A", "out"),
    ("A_in", "A", "in"),
    ("x_out", "x", "out"),
    ("x_in", "x", "in"),
    ("C_h", "C", "h"),
    ("R_h", "R", "h"),
    ("k_s", "k", "s"),
    ("x_f", "x", "f"),
    ("p0", "p", "0"),
    ("V0", "V", "0"),
    ("t10", "t", "10"),
    ("t63", "t", "63"),
    ("t90", "t", "90"),
]
_SUB_MAP = {t: (b, s) for t, b, s in SUB_TOKENS}
_SUB_RE = re.compile(
    r"(?<![A-Za-z0-9_])(?:"
    + "|".join(re.escape(t) for t, _, _ in sorted(SUB_TOKENS, key=lambda x: -len(x[0])))
    + r")(?![A-Za-z0-9_])"
)


def set_font(run, name="Times New Roman", size=10.5, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def en_text(s):
    return (str(s).replace("mm2", "mm\u00b2")
            .replace("tau", "\u03c4")
            .replace("Delta", "\u0394")
            .replace("pV^n", "pV\u207f")
            .replace("3.1x", "3.1\u00d7")
            .replace("6x", "6\u00d7"))


def _replace_run(paragraph, run, pieces):
    rpr = run._r.find(qn("w:rPr"))
    parent = run._r.getparent()
    index = parent.index(run._r)
    parent.remove(run._r)
    for mode, s in pieces:
        if not s:
            continue
        r = OxmlElement("w:r")
        if rpr is not None:
            r.append(deepcopy(rpr))
        t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = s
        r.append(t)
        if mode:
            rp = r.find(qn("w:rPr"))
            if rp is None:
                rp = OxmlElement("w:rPr")
                r.insert(0, rp)
            for va in rp.findall(qn("w:vertAlign")):
                rp.remove(va)
            va = OxmlElement("w:vertAlign")
            va.set(qn("w:val"), "subscript" if mode == "sub" else "superscript")
            rp.append(va)
        parent.insert(index, r)
        index += 1


def _sub_paragraph(p):
    for run in list(p.runs):
        text = run.text
        if not text or not _SUB_RE.search(text):
            continue
        pieces = []
        pos = 0
        for m in _SUB_RE.finditer(text):
            if m.start() > pos:
                pieces.append((None, text[pos:m.start()]))
            base, sub = _SUB_MAP[m.group(0)]
            if base:
                pieces.append((None, base))
            pieces.append(("sub", sub))
            pos = m.end()
        if pos < len(text):
            pieces.append((None, text[pos:]))
        _replace_run(p, run, pieces)


def _sup_paragraph(p):
    for run in list(p.runs):
        text = run.text
        if not text or "\u00b2" not in text:
            continue
        pieces = []
        pos = 0
        for i, ch in enumerate(text):
            if ch == "\u00b2":
                if i > pos:
                    pieces.append((None, text[pos:i]))
                pieces.append(("sup", "2"))
                pos = i + 1
        if pos < len(text):
            pieces.append((None, text[pos:]))
        _replace_run(p, run, pieces)


def subscriptify(doc):
    paras = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                paras.extend(cell.paragraphs)
    for p in paras:
        _sub_paragraph(p)
        _sup_paragraph(p)


def H(doc, text, level):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    for r in p.runs:
        set_font(r, "Arial", {1: 15, 2: 12.5, 3: 11}.get(level, 11), True,
                 RGBColor(0, 0, 0))
    return p


def P(doc, text, size=10.5, bold=False, align=None, indent=True):
    text = en_text(text)
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.4
    p.paragraph_format.space_after = Pt(4)
    if indent:
        p.paragraph_format.first_line_indent = Cm(0.75)
    r = p.add_run(text)
    set_font(r, size=size, bold=bold)
    if align is not None:
        p.alignment = align
    return p


def EQ(doc, text):
    text = en_text(text)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.4
    pos = 0
    for m in re.finditer(r"\^\{([^}]*)\}", text):
        if m.start() > pos:
            r = p.add_run(text[pos:m.start()])
            set_font(r, "Times New Roman", 10.5, True)
        r = p.add_run(m.group(1))
        set_font(r, "Times New Roman", 10.5, True)
        r.font.superscript = True
        pos = m.end()
    if pos < len(text):
        r = p.add_run(text[pos:])
        set_font(r, "Times New Roman", 10.5, True)
    return p


def B(doc, text):
    text = en_text(text)
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.line_spacing = 1.35
    set_font(p.add_run(text))
    return p


def CAP(doc, text):
    text = en_text(text)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    set_font(r, "Times New Roman", 9, False, RGBColor(0x55, 0x55, 0x55))
    return p


def PIC(doc, path, width, caption):
    if os.path.exists(path):
        doc.add_picture(path, width=Cm(width))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        CAP(doc, caption)


def TABLE(doc, header, rows, widths=None, font=9, caption=None):
    if caption:
        CAP(doc, caption)
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, s in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(en_text(s))
        set_font(r, "Arial", font, True)
    for row in rows:
        cells = t.add_row().cells
        for i, s in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(en_text(s))
            set_font(r, "Times New Roman", font)
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Cm(w)
    tr = t.rows[0]._tr
    trPr = tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def PAGEBREAK(doc):
    doc.add_page_break()


def add_page_number_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Page ")
    set_font(r, "Times New Roman", 9)
    run = p.add_run()
    fld1 = OxmlElement("w:fldChar"); fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = "PAGE"
    fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1); run._r.append(instr); run._r.append(fld2)
    set_font(run, "Times New Roman", 9)


# ---------------------------------------------------------------------------
# document
# ---------------------------------------------------------------------------

def main():
    doc = Document()
    doc.core_properties.author = "Zihao Huang"
    doc.core_properties.last_modified_by = "Zihao Huang"
    doc.core_properties.title = ("Bubble Compliance in a Syringe Hydraulic Transmission: "
                                 "Static Stiffness, Dynamic Response, and Variable Stiffness")

    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.4)
    sec.bottom_margin = Cm(2.4)
    sec.left_margin = Cm(2.5)
    sec.right_margin = Cm(2.5)
    sec.header_distance = Cm(1.1)
    sec.footer_distance = Cm(1.1)
    sec.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.4
    normal.paragraph_format.space_after = Pt(4)

    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hr = hp.add_run("Bubble compliance in syringe hydraulic transmission")
    set_font(hr, "Times New Roman", 8.5, False, RGBColor(0x77, 0x77, 0x77))
    add_page_number_footer(sec)

    # cover
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Bubble Compliance in a Syringe Hydraulic Transmission")
    set_font(r, "Arial", 20, True, RGBColor(0x8B, 0x1A, 0x1A))
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Static Stiffness, Dynamic Response, and Variable Stiffness")
    set_font(r, "Arial", 14, True, RGBColor(0x8B, 0x1A, 0x1A))
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("An experimental study with 1, 2, and 5 mL air volumes")
    set_font(r, "Arial", 11.5, False)
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Technical Report")
    set_font(r, "Arial", 12, True)
    for line in ["Zihao Huang", "October 2026",
                 "All data are measured; no simulated data are used."]:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(line); set_font(r, "Times New Roman", 11)
    PAGEBREAK(doc)

    # abstract
    H(doc, "Abstract", 1)
    P(doc, "This report studies how a trapped air bubble changes the static stiffness, dynamic response, and "
           "variable-stiffness capability of a small syringe hydraulic transmission. The circuit uses a 10 mL "
           "input syringe, a 20 mL output syringe, a silicone tube, and an output extension spring. Known air "
           "volumes of 1, 2, and 5 mL are injected to set the effective compliance. Static tests apply known "
           "water weights (0.3, 0.55, and 0.85 kg); dynamic tests use a fast input step and 240 fps video. "
           "All results are calculated from measured CSV data with 2-4 repeats per condition.", indent=False)
    P(doc, "The static stiffness K_air decreases from 6.82±1.60 to 3.41±0.13 and 1.38±0.02 N/mm for 1, 2, and "
           "5 mL. The 1/K_air-V relation is linear, giving C = 6.95 N·mm/mL, a baseline air volume V0 = 0.03 mL, "
           "and a polytropic index n = 1.00, consistent with an isothermal static process. The time constant tau "
           "increases linearly with air volume: 0.113±0.001, 0.181±0.013, and 0.409±0.029 s, with "
           "tau = 0.035 + 0.074V (three-point fit, R² = 0.9995). The soft-joint stiffness K_eff is 7.44±0.25, "
           "4.57±0.09, and 2.40±0.05 N/mm, a 3.1x tuning range over 1-5 mL, and agrees with K_eff = k_s + K_air "
           "within 5%. The results show that a trapped bubble can act as a passive, valve-free variable-stiffness "
           "element, and that the response time can be used as a bubble index. The main limitations are the "
           "resolution at 1 mL, the ramp input, a single output area, and the absence of pressure and temperature "
           "sensors.", indent=False)
    p = doc.add_paragraph()
    r = p.add_run("Keywords: "); set_font(r, "Times New Roman", 10.5, True)
    r = p.add_run("hydraulic transmission; bubble compliance; air spring; variable stiffness; "
                  "time constant; soft robotics")
    set_font(r, "Times New Roman", 10.5)
    PAGEBREAK(doc)

    # 1 introduction
    H(doc, "1  Introduction", 1)
    H(doc, "1.1  Background", 2)
    P(doc, "Hydraulic transmission converts force and displacement through a confined fluid and offers high "
           "power density, flexible routing, and direct linear actuation [4]. In small syringe-based systems, "
           "the fluid volume is small, the tubing is narrow, and the number of seals is large, so air is easily "
           "trapped as bubbles in the syringe or tube.")
    P(doc, "Bubbles are usually undesirable: they reduce the effective stiffness, delay the output, and cause "
           "pressure fluctuations, noise, and energy loss [4]. From another perspective, a trapped bubble is a "
           "compressible air spring whose compliance is proportional to its volume [4]. If the bubble volume is "
           "controlled, the bubble can be turned from a defect into a passive, valve-free compliance element.")
    H(doc, "1.2  Soft robotics and variable stiffness", 2)
    P(doc, "Soft robots are built from elastomers, fibre-reinforced layers, or flexible chambers and offer high "
           "compliance and safe human interaction [1,2,5]. Their low stiffness, however, limits load capacity and "
           "positioning accuracy. Variable-stiffness actuators switch between soft and stiff states in the same "
           "device, so that the robot can be compliant during contact and stiff when a load must be carried [1,2,5].")
    P(doc, "Existing variable-stiffness approaches include granular jamming, layer jamming, low-melting-point "
           "materials, antagonistic actuation, and additional springs [3,6]. These approaches often require extra "
           "structures, valves, or energy input. A syringe hydraulic circuit with a controlled bubble offers a "
           "simpler route: the air volume is a passive, valve-free, and repeatable stiffness knob.")
    H(doc, "1.3  Bubble detection", 2)
    P(doc, "Bubbles are also a target for fluid-line health monitoring [4]. A bubble slows the response and "
           "degrades displacement transfer. If the bubble volume can be inferred from the output response, "
           "incipient air ingestion or leakage can be detected without a dedicated bubble sensor. The time "
           "constant is a candidate bubble index because it grows with the trapped air volume.")
    H(doc, "1.4  Scope and questions", 2)
    P(doc, "This work does not use a load cell; it uses known water weights, an output spring, and displacement "
           "measurement. The study addresses four questions:")
    B(doc, "How does the bubble volume V change the static equivalent stiffness K_air?")
    B(doc, "How does V change the step-response time constant tau, and is tau proportional to V?")
    B(doc, "Can the output spring and the bubble be combined into a usable variable-stiffness joint, and how "
           "large is the stiffness range?")
    B(doc, "Are the static stiffness, dynamic response, and spring additivity self-consistent, and can the "
           "result be used for bubble detection?")
    H(doc, "1.5  Main results", 2)
    P(doc, "Three experiments were performed. The static stiffness K_air falls from 6.82 to 1.38 N/mm as V "
           "increases from 1 to 5 mL, and 1/K_air is linear in V. The time constant tau rises from 0.113 to "
           "0.409 s and is linear in V (three-point fit, R² = 0.9995). The soft-joint stiffness K_eff falls from "
           "7.44 to 2.40 N/mm, a 3.1x range, and agrees with k_s + K_air within 5%. The three data sets are "
           "mutually consistent.")
    H(doc, "1.6  Structure", 2)
    P(doc, "Section 2 presents the air-spring, static, dynamic, and volume-balance models. Section 3 describes "
           "the setup, calibration, bubble injection, and the three experiments. Section 4 reports the measured "
           "results and the cross-checks. Section 5 discusses the engineering implications, bubble detection, "
           "and limitations. Section 6 concludes.")
    PAGEBREAK(doc)

    # 2 theory
    H(doc, "2  Theory", 1)
    H(doc, "2.1  Bubble as an air spring", 2)
    P(doc, "Let p0 and V be the initial absolute pressure and volume of the trapped air. The compression follows "
           "the polytropic relation pV^n = constant [4], where n is the polytropic index: n is about 1 for a "
           "slow, isothermal process and about 1.4 for a fast, adiabatic process. Linearising about the "
           "equilibrium gives")
    EQ(doc, "dp = - n p0 dV / V")
    P(doc, "For an output piston of area A_out and a small displacement x, the air volume changes by "
           "dV = -A_out x. The pressure change therefore produces an equivalent stiffness")
    EQ(doc, "K_air = n p0 A_out^{2} / V")
    P(doc, "Equation (2) shows that K_air is inversely proportional to V: more air makes the system softer. "
           "For 5 mL, A_out = 262 mm2, and n = 1, K_air is about 1.4 N/mm, the same order as the output spring "
           "(about 1 N/mm); the bubble is therefore the main compliance source.")
    H(doc, "2.2  Static stiffness and 1/K linearity", 2)
    P(doc, "With the output spring removed, the output is loaded only by the water weight and the air pressure. "
           "Let V0 be the residual air volume and V the injected volume. The total air volume is V0 + V, and")
    EQ(doc, "1 / K_air = (V0 + V) / C,   C = n p0 A_out^{2}")
    P(doc, "Thus 1/K_air is linear in V, with slope 1/C and intercept V0/C. When a spring of stiffness k_s is "
           "installed in parallel with the bubble, the static equivalent stiffness is")
    EQ(doc, "K_eff = k_s + K_air")
    P(doc, "Equation (4) is the criterion for the soft-joint experiment: the measured K_eff minus K_air should "
           "equal the calibrated spring stiffness.")
    H(doc, "2.3  Dynamic response and time constant", 2)
    P(doc, "The hydraulic circuit is treated as a first-order system: the bubble provides the hydraulic "
           "capacitance C_h, and the tubing and syringes provide the hydraulic resistance R_h. For a small "
           "disturbance, the output displacement follows")
    EQ(doc, "x(t) = x_f [1 - exp(-t / tau)],   tau = R_h C_h")
    P(doc, "where x_f is the final displacement. The hydraulic capacitance of the bubble is")
    EQ(doc, "C_h = V / (n p0)")
    P(doc, "and therefore")
    EQ(doc, "tau = R_h V / (n p0) = (R_h / (n p0)) (V0 + V)")
    P(doc, "Equation (7) predicts that tau is linear in V with slope R_h/(n p0). A measured linear tau-V "
           "relation therefore supports the assumptions that bubble compliance dominates and that the hydraulic "
           "resistance is approximately constant.")
    H(doc, "2.4  Volume balance and final displacement", 2)
    P(doc, "Let A_in and x_in be the input area and displacement, and A_out and x_f the output area and final "
           "displacement. The input volume pushes the output and compresses the air:")
    EQ(doc, "A_in x_in = A_out x_f + V dp / p0")
    P(doc, "The output spring force balances the pressure, dp = k_s x_f / A_out. Substitution gives")
    EQ(doc, "x_f = A_in x_in / [A_out + V k_s / (p0 A_out)]")
    P(doc, "The term V k_s/(p0 A_out) in the denominator acts as an equivalent bubble volume: part of the input "
           "liquid is used to compress the air, so the effective volume that drives the output piston is reduced "
           "and x_f decreases as V increases. This explains why the 5 mL case has a smaller output displacement "
           "than the 2 mL case.")
    H(doc, "2.5  Link between static and dynamic behaviour", 2)
    P(doc, "The static stiffness K_air = n p0 A_out²/V and the dynamic time constant tau = R_h V/(n p0) both "
           "follow from the same hydraulic capacitance C_h = V/(n p0). The static 1/K slope gives C and V0; the "
           "dynamic tau slope gives R_h/(n p0). The two data sets can therefore be used to check each other and "
           "to estimate the order of magnitude of R_h.")
    H(doc, "2.6  Main parameters", 2)
    TABLE(doc, ["Parameter", "Symbol", "Value", "Unit / note"],
          [["Input syringe area", "A_in", "160", "mm2, nominal for 10 mL"],
           ["Output syringe area", "A_out", "262", "mm2, from scale spacing and air-spring fit"],
           ["Area ratio", "A_out/A_in", "1.64", "-"],
           ["Atmospheric pressure", "p0", "0.1013", "N/mm2 (101.3 kPa)"],
           ["Polytropic index", "n", "1.00", "isothermal static fit"],
           ["Output spring stiffness", "k_s", "1.02", "N/mm, water calibration (3 repeats)"],
           ["Static compliance parameter", "C", "6.95", "N·mm/mL"],
           ["Baseline air volume", "V0", "0.03", "mL"],
           ["Dynamic slope", "m", "0.074", "s/mL"]],
          widths=[5.2, 2.6, 2.6, 6.6], font=9,
          caption="Table 1. Main parameters and units.")
    PAGEBREAK(doc)

    # 3 methods
    H(doc, "3  Experimental setup and methods", 1)
    H(doc, "3.1  Setup", 2)
    P(doc, "The setup consists of mechanical, hydraulic, drive, and measurement parts. The mechanical part "
           "includes a stepper motor, a lead-screw linear stage, an input-syringe holder, an output-syringe "
           "holder, an output extension spring, and a fixed end. The hydraulic part is a 10 mL input syringe, a "
           "20 mL output syringe, and a silicone tube. The drive part is an Arduino Mega2560, a TB6600 driver, "
           "and a 12 V supply. The measurement part is a set of known water weights (0.3, 0.55, and 0.85 kg), a "
           "caliper, a 0.2 mm printed scale, and a 240 fps phone camera. No load cell is used; the force is set "
           "by the known water weight and the calibrated spring.")
    PIC(doc, os.path.join(DATA, "setup_schematic_en.png"), 15.5,
        "Figure 1. Schematic of the syringe hydraulic bubble-compliance setup.")
    P(doc, "Figure 2 shows the hydraulic circuit and displacement measurement. The output syringe and extension "
           "spring are on the left, the input syringe and lead screw are on the right, and the phone on its stand "
           "records the output at 240 fps. The water load and the pulley are shown schematically in Figure 1.")
    PIC(doc, os.path.join(DATA, "setup_photo_en.jpg"), 14.0,
        "Figure 2. Hydraulic circuit and displacement measurement (output syringe and extension spring, input "
        "syringe and lead screw, 0.2 mm scale, phone stand and 240 fps phone).")
    TABLE(doc, ["Component", "Specification", "Role"],
          [["Input syringe", "10 mL, A_in = 160 mm2", "hydraulic input"],
           ["Output syringe", "20 mL, A_out = 262 mm2", "hydraulic output"],
           ["Silicone tube", "4 mm inner diameter", "connects the syringes"],
           ["Output spring", "k_s = 1.02 N/mm", "load, return, soft joint"],
           ["Water weights", "0.3, 0.55, 0.85 kg", "known static force 2.94/5.39/8.34 N"],
           ["Linear stage", "2 mm lead, 3200 steps/rev", "input displacement"],
           ["Driver", "TB6600, 0.5 A, 1/16 microstep", "stepper drive"],
           ["Controller", "Arduino Mega2560", "pulses and serial"],
           ["Camera", "240 fps phone + 0.2 mm scale", "dynamic displacement"],
           ["Caliper", "0.02 mm resolution", "static displacement"],
           ["Printed parts", "syringe clamps, rod clamp, brackets", "fixture and force transfer"]],
          widths=[3.4, 7.8, 5.2], font=8.5,
          caption="Table 2. Main hardware.")
    H(doc, "3.2  Area and calibration", 2)
    P(doc, "The output area is not taken from the nominal 20 mL value of 287 mm2. The 20 mL syringe has 0.5 mL "
           "per small division and a measured spacing of about 1.9 mm, giving A_out = 500/1.9 = 263 mm2. The "
           "static 1/K fit gives C = 6.95 N·mm/mL and, with n = 1, A_out = 262 mm2. The two estimates agree. "
           "The input area is the nominal 10 mL value, A_in = 160 mm2; the area ratio is 1.64.")
    P(doc, "The output spring was calibrated independently by hanging known water weights. Three repeat "
           "F-Deltax slopes were 1.018, 1.013, and 1.038 N/mm; the pooled fit gives k_s = 1.02 N/mm, "
           "b = -0.24 N, and R² = 0.999 (Figure 3). The bottle and string weight affect only the intercept, not "
           "the slope.")
    PIC(doc, os.path.join(DATA, "spring_calib_en.png"), 11.0,
        "Figure 3. Output-spring calibration: force versus extension for three repeats.")
    H(doc, "3.3  Bubble injection and baseline", 2)
    P(doc, "Before each test, the syringes were repeatedly pushed, tapped, and tilted to remove visible bubbles; "
           "this is the 0 mL baseline. A three-way valve and a 1 mL syringe were used to inject 1, 2, or 5 mL of "
           "air; when a three-way valve was not available, the output tube was removed, the output rod was "
           "pulled to draw in the set volume, and the tube was reconnected. The injected volume was read from "
           "the syringe scale with an uncertainty of about ±0.05 mL. The residual volume V0 was not measured "
           "directly; it was obtained from the intercept of the 1/K-V fit (V0 = 0.03 mL). The 0 mL and 0.2 mL "
           "cases were too stiff for the available displacement resolution and were not used in the quantitative "
           "fit.")
    H(doc, "3.4  Static stiffness experiment", 2)
    P(doc, "The output spring was removed. A clamp locked the input rod so that the input did not move and the "
           "total air volume stayed constant. The output rod was connected by a thin string over a smooth bar to "
           "a water bottle, converting the vertical weight into a horizontal force along the output axis. The "
           "0.3, 0.55, and 0.85 kg weights were applied in turn; after 5-10 s the output displacement x_out was "
           "measured with the caliper. Each bubble volume was tested three times.")
    H(doc, "3.5  Step-response experiment", 2)
    P(doc, "The output spring was reinstalled and the input lock was released. The stepper drove the input from "
           "the same starting point in each run. After 1-2 s of baseline, a fast step was commanded and the output "
           "rod and scale were filmed from above at 240 fps until the output had been flat for at least 2 s. The "
           "video time was computed as t = frame/240 s; the slow-motion time shown by the player was not used.")
    P(doc, "For 2 and 5 mL, the input step was 2 mm at 10 mm/s (0.2 s ramp). For 1 mL, whose time constant is "
           "smaller, the step was 0.5 mm at 16 mm/s (0.031 s ramp) and the output was about 0.27 mm, read with "
           "the 0.2 mm scale. The time constant was defined at 63.2% of the total change. To reduce the effect of "
           "the input ramp, the script used the 10% and 63% crossings: tau = (t63 - t10)/0.8946.")
    P(doc, "Three early 5 mL runs were excluded: 0.109 and 0.205 s, where the output reached its final value too "
           "early and the bubble position or input state was abnormal, and 0.425 s, which used a different "
           "4 mm/8 mm/s protocol. One early 1 mL run (0.118 s) was excluded because the rising part was sampled "
           "too sparsely; after additional sampling, 0.113 and 0.112 s were used. All exclusions are documented.")
    H(doc, "3.6  Soft-joint experiment", 2)
    P(doc, "The output spring was installed and the input was locked, so the total air volume was fixed. The "
           "output spring acted as the elastic element of a soft joint. For each bubble volume, the 0.3, 0.55, "
           "and 0.85 kg weights were applied and the output displacement was measured; K_eff was obtained from "
           "the F-x slope. Each bubble volume was tested three times. K_eff was compared with k_s + K_air, and "
           "the stiffness range over 1-5 mL was calculated.")
    H(doc, "3.7  Data processing", 2)
    P(doc, "The raw data are CSV files. For the static tests the columns are air volume, load, displacement, and "
           "repeat; for the dynamic tests they are time and output displacement. Python scripts group the data by "
           "bubble volume and repeat, fit each three-point set to obtain K, compute the mean, standard deviation, "
           "and R², extract tau at 63.2%, and fit 1/K-V and tau-V. All figures and tables in this report are "
           "generated from the same measured CSV files; no simulated data are used.")
    H(doc, "3.8  Uncertainty", 2)
    P(doc, "The main uncertainty sources are the water mass (about ±1%; the bottle and string weight are treated "
           "as an intercept), the bubble volume (±0.05 mL), the static displacement (caliper ±0.02 mm; for 1 mL "
           "the 0.3 kg displacement is only 0.4-0.5 mm, so the relative error can reach 10-20%), the dynamic "
           "displacement (video ±0.05-0.1 mm; the 1 mL output is only 0.27 mm), the output area (±3%), the "
           "spring stiffness (calibrated, ±1%), and temperature and bubble position (not actively controlled). "
           "The standard deviation of the repeats is used as the main error estimate, and the raw points are "
           "retained in the figures.")
    PAGEBREAK(doc)

    # 4 results
    H(doc, "4  Results", 1)
    H(doc, "4.1  Static bubble stiffness", 2)
    P(doc, "Table 3 summarises the static stiffness. K_air is 6.82±1.60 N/mm at 1 mL (relative error about 23%), "
           "3.41±0.13 N/mm at 2 mL (3.9%), and 1.38±0.02 N/mm at 5 mL (1.1%). K_air decreases with increasing "
           "bubble volume, following the 1/V law.")
    TABLE(doc, ["V (mL)", "K_air (N/mm)", "SD (N/mm)", "Relative error", "n"],
          [["1", "6.82", "1.60", "23.4%", "3"],
           ["2", "3.41", "0.13", "3.9%", "3"],
           ["5", "1.38", "0.02", "1.1%", "3"]],
          widths=[2.6, 3.6, 3.0, 3.2, 1.6], font=9,
          caption="Table 3. Static bubble stiffness.")
    P(doc, "Figure 4 shows K_air and 1/K_air against V. Grey circles are individual repeat fits, blue squares are "
           "means ± SD, and the red line is the model C/(V0 + V). The 1 mL points scatter more, but the 2 and 5 mL "
           "points are stable; the linear 1/K-V fit gives R² = 1.0000.")
    PIC(doc, os.path.join(DATA, "static_stiffness_en.png"), 15.5,
        "Figure 4. Static stiffness and inverse stiffness versus air volume (individual points, mean ± SD, "
        "and model).")
    P(doc, "The linear fit is")
    EQ(doc, "1/K_air = 0.0038 + 0.1439 V   (R^{2} = 1.0000)")
    P(doc, "giving C = 6.95 N·mm/mL and V0 = 0.03 mL. Using A_out = 262 mm2,")
    EQ(doc, "n = C x 1000 / (p0 A_out^{2}) = 1.00")
    P(doc, "which is consistent with an isothermal static process; the theoretical C = n p0 A_out²/1000 = "
           "6.95 N·mm/mL differs from the measured value by less than 1%. The output area, the bubble volume, "
           "and the air-spring model are therefore mutually consistent.")
    P(doc, "The 1 mL scatter needs a separate comment. Its 0.3 kg displacement is only 0.4-0.5 mm; the caliper "
           "resolution and the initial-position error are of order 0.1 mm, which can produce a slope error of "
           "10-20%. The three individual K values are 6.51, 5.39, and 8.54 N/mm. The 1/K fit is therefore "
           "anchored mainly by the 2 and 5 mL data.")
    P(doc, "The 1/K-V fit uses only the three mean points at 1, 2, and 5 mL. A three-point fit has only one "
           "degree of freedom, so R² is easily close to 1; it shows that the three points are nearly collinear "
           "but is not a test of statistical significance. The reliability of the model comes from the "
           "repeatability of the measurements and from the independent static, dynamic, and soft-joint checks.")
    H(doc, "4.2  Step response and time constant", 2)
    P(doc, "Table 4 summarises the step response. tau is 0.113±0.001 s at 1 mL (n = 2), 0.181±0.013 s at 2 mL "
           "(n = 4), and 0.409±0.029 s at 5 mL (n = 4). tau increases with bubble volume, as expected from "
           "C_h proportional to V.")
    TABLE(doc, ["V (mL)", "tau (s)", "SD (s)", "Relative error", "n"],
          [["1", "0.113", "0.001", "0.7%", "2"],
           ["2", "0.181", "0.013", "7.1%", "4"],
           ["5", "0.409", "0.029", "7.0%", "4"]],
          widths=[2.6, 3.6, 3.0, 3.2, 1.6], font=9,
          caption="Table 4. Step-response time constant.")
    P(doc, "Figure 5 shows tau against V. The three points are approximately collinear, and the ratio "
           "tau(5 mL)/tau(2 mL) = 2.26 is close to the volume ratio 2.5.")
    PIC(doc, os.path.join(DATA, "step_tau_en.png"), 13.0,
        "Figure 5. Time constant versus air volume (individual points, mean ± SD, and linear fit).")
    P(doc, "The linear fit is")
    EQ(doc, "tau = 0.035 + 0.074 V   (R^{2} = 0.9995)")
    P(doc, "The predicted values are 0.109, 0.183, and 0.405 s for 1, 2, and 5 mL; the measured values are "
           "0.113, 0.181, and 0.409 s, within about 3%. The intercept of 0.035 s corresponds to residual bubble "
           "and line compliance. The R² of this three-point fit is also only a measure of collinearity, not a "
           "significance test.")
    P(doc, "Figure 6 shows the normalized step responses, one representative curve per bubble volume, aligned at "
           "the 10% rise time t10. The 1 mL response is the fastest and the 5 mL response the slowest; all curves "
           "reach 1, indicating no significant leakage or lost steps. The inset zooms the first 0.4 s of the "
           "1 mL response.")
    PIC(doc, os.path.join(DATA, "step_examples_en.png"), 13.0,
        "Figure 6. Normalized step responses for 1, 2, and 5 mL (representative curve per volume; time is "
        "relative to t10; inset: 1 mL zoom).")
    H(doc, "4.3  Soft-joint variable stiffness", 2)
    P(doc, "Table 5 gives the soft-joint stiffness. K_eff is 7.44±0.25 N/mm at 1 mL, 4.57±0.09 N/mm at 2 mL, and "
           "2.40±0.05 N/mm at 5 mL; the stiffness range over 1-5 mL is 7.44/2.40 = 3.1x.")
    TABLE(doc, ["V (mL)", "K_eff (N/mm)", "SD (N/mm)", "Relative error", "n"],
          [["1", "7.44", "0.25", "3.3%", "3"],
           ["2", "4.57", "0.09", "1.9%", "3"],
           ["5", "2.40", "0.05", "2.1%", "3"]],
          widths=[2.6, 3.6, 3.0, 3.2, 1.6], font=9,
          caption="Table 5. Soft-joint equivalent stiffness.")
    P(doc, "Figure 7 shows K_eff against V. Grey circles are individual repeats, blue squares are means ± SD, and "
           "the red line is the model k_s + C/(V0 + V) with k_s = 1.02 N/mm.")
    PIC(doc, os.path.join(DATA, "demo_stiffness_en.png"), 13.0,
        "Figure 7. Soft-joint stiffness versus air volume (individual points, mean ± SD, and model).")
    P(doc, "Table 6 compares K_eff with the prediction k_s + K_air. The deviations are -5.1%, +3.0%, and -0.3%, "
           "all within about 5%. The measured spring stiffness is k_s = 1.02±0.01 N/mm; the values back-calculated "
           "from K_eff - K_air are 0.62, 1.16, and 1.01 N/mm (mean about 0.93 N/mm), close to the calibrated "
           "value. The -5.1% deviation at 1 mL comes mainly from the scatter of the 1 mL static K_air; the +3.0% "
           "at 2 mL may come from a slight underestimate of the static K_air. Considering the 1 mL reading error, "
           "the spring additivity holds.")
    TABLE(doc, ["V (mL)", "K_air", "k_s + K_air", "K_eff measured", "Deviation"],
          [["1", "6.82", "7.84", "7.44", "-5.1%"],
           ["2", "3.41", "4.43", "4.57", "+3.0%"],
           ["5", "1.38", "2.40", "2.40", "-0.3%"]],
          widths=[2.4, 2.9, 3.6, 3.6, 2.8], font=9,
          caption="Table 6. Spring additivity (k_s = 1.02 N/mm). k_s is the output-spring stiffness and K_air "
                  "is the bubble stiffness.")
    H(doc, "4.4  Cross-validation", 2)
    P(doc, "Four independent checks are used:")
    B(doc, "Static theory: the measured C is 6.95 N·mm/mL, the theoretical value is 6.95 N·mm/mL, and n = 1.00.")
    B(doc, "Spring additivity: K_eff agrees with k_s + K_air within about 5% (k_s = 1.02 N/mm, water calibration).")
    B(doc, "Static-dynamic consistency: the static data give C_h proportional to V, and the dynamic data give "
           "tau proportional to V; the two slopes are of the same order.")
    B(doc, "Volume balance: the final displacements for 5 mL/2 mm, 2 mL/2 mm, and 1 mL/0.5 mm are 0.70, 0.95, "
           "and 0.27 mm, consistent with Eq. (9).")
    TABLE(doc, ["Check", "Expected", "Measured", "Conclusion"],
          [["Static C", "n p0 A_out²/1000 = 6.95", "6.95", "difference <1%"],
           ["K_eff additivity", "k_s + K_air", "7.44/4.57/2.40", "within about 5%"],
           ["tau-V linearity", "tau proportional to V", "R² = 0.9995", "three points collinear (R² for reference)"],
           ["Final displacement", "Eq. (9)", "0.70/0.95/0.27 mm", "consistent"]],
          widths=[3.8, 5.0, 4.4, 4.2], font=8.5,
          caption="Table 7. Cross-validation summary.")
    P(doc, "The validation uses analytical models and independent experiments, not simulated data. All numbers "
           "in this report come from measured CSV files.")
    PAGEBREAK(doc)

    # 5 discussion
    H(doc, "5  Discussion", 1)
    H(doc, "5.1  Bubble as an air spring", 2)
    P(doc, "The static data give a clean 1/K_air-V line, and C differs from the theoretical value by less than "
           "1%, with n = 1.00 [4]. Under slow, well-exchanged conditions, the trapped air therefore behaves as an "
           "isothermal air spring, and the bubble volume is the main compliance source. The compliance of the "
           "tubing and syringes appears only in the small V0 of 0.03 mL. This provides a quantitative basis for "
           "using the bubble as an adjustable spring: changing the injected volume sets the stiffness according "
           "to 1/V.")
    P(doc, "The 1 mL scatter is a reminder that the air-spring model is limited by the gauge and by static "
           "friction at very small volume and displacement. The 0.3 kg displacement at 1 mL is only 0.4-0.5 mm; "
           "the 0.02 mm caliper resolution and a 0.05 mm initial-position error are enough to cause more than 10% "
           "slope error. For high-stiffness applications, a finer displacement sensor or a longer loading stroke "
           "should be used instead of a 0.5 mm scale.")
    H(doc, "5.2  Static-dynamic consistency", 2)
    P(doc, "Static and dynamic data point to the same C_h = V/(n p0). The static 1/K slope gives C = 6.95 "
           "N·mm/mL; the dynamic tau slope gives 0.074 s/mL. The two cannot be equated directly because the units "
           "differ, but they can be used to check each other: from tau = R_h C_h, the tau-V slope should be "
           "proportional to R_h/(n p0); from the static data, C = n p0 A_out². Combining the two gives the order "
           "of magnitude of R_h and shows that it is approximately independent of the bubble volume. The 1 mL tau "
           "is slightly higher, probably because of the input ramp or the bubble position.")
    P(doc, "The dynamic tau was measured with 2 mm at 10 mm/s (0.2 s ramp) or 0.5 mm at 16 mm/s (0.031 s ramp), "
           "so it is an effective time constant that includes the input shape. The 2 mL tau of 0.181 s is close "
           "to the 0.2 s ramp, so the absolute tau may be slightly overestimated. The 2 and 5 mL data use the same "
           "protocol, so their ratio of 2.26 can still be attributed to the air volume. The 1 mL run uses a faster "
           "ramp, and its tau of 0.113 s differs by only 3% from the 0.109 s extrapolated from the 2/5 mL line, "
           "so the protocol difference does not break the trend.")
    H(doc, "5.3  Variable stiffness for soft robotics", 2)
    P(doc, "The soft-joint experiment shows that the output spring and the bubble in parallel reduce K_eff from "
           "7.44 to 2.40 N/mm, a range of about 3.1x over 1-5 mL. Extrapolating to 0.5 mL gives K_eff of about "
           "14 N/mm and a range of about 6x (not measured). This range is relevant to small soft joints, grippers, "
           "and wearable devices [1,2]: with one hydraulic source, the same actuator can switch between compliant "
           "contact and load-bearing positioning by changing the air volume.")
    P(doc, "The cost is a slower response: as K_eff falls from 7.44 to 2.40 N/mm, tau rises from 0.113 to 0.409 s. "
           "For fast tasks, the air volume should be small; for compliant, impact-tolerant tasks, it can be "
           "larger. This soft-slow trade-off is inherent to bubble compliance and must be considered in design. "
           "Compared with granular jamming and low-melting-point materials [3,6], the bubble approach is "
           "structurally simple, valve-free, and requires no extra energy input; its limitations are the "
           "stiffness range set by the bubble volume, the slower response, and long-term leakage and bubble "
           "coalescence.")
    H(doc, "5.4  Bubble detection", 2)
    P(doc, "The linear tau-V relation provides a basis for bubble detection [4]. Under a fixed input protocol, a "
           "single step response gives tau and hence the air volume from tau = 0.035 + 0.074V. Using the standard "
           "deviation of the repeats (3 sigma), the resolution is about 0.5 mL near 2 mL and about 1.2 mL near "
           "5 mL; the 1 mL displacement is too small and the number of repeats too low to set a detection "
           "threshold. The method is therefore suitable for detecting obvious air ingestion or leakage, not very "
           "small bubbles. In practice, the bubble position (syringe versus tube) changes the dynamic response, "
           "temperature affects p0 and n, and the input ramp mixes into tau; the tau-V curve should therefore be "
           "calibrated in situ with the same fixture and protocol.")
    H(doc, "5.5  Limitations", 2)
    B(doc, "The 1 mL static displacement is only 0.4-0.5 mm and K_air has a 23% relative error; the 1 mL step "
           "response has only two repeats.")
    B(doc, "The dynamic input is a 0.2 s or 0.031 s ramp, not an ideal step; the 2 mL tau may be overestimated by "
           "the input shape.")
    B(doc, "Bubble position, coalescence, and leakage were not tightly controlled; three early 5 mL runs were "
           "excluded for this reason.")
    B(doc, "Only 1, 2, and 5 mL were measured; 0.5 mL and below are extrapolated.")
    B(doc, "Only one output area (A_out = 262 mm2) was used; the A_out² scaling is supported indirectly by the "
           "scale spacing and n = 1, but has not been independently tested with 5 and 10 mL output syringes.")
    B(doc, "No pressure or temperature sensor was used; n = 1 is a static inference and the dynamic n may differ.")
    B(doc, "The spring stiffness was calibrated by the water method, but the calibration pose may differ from the "
           "installed pose; the k_s uncertainty enters K_eff.")
    B(doc, "The video spatial resolution is about 0.05-0.1 mm, which is marginal for the 0.27 mm 1 mL output.")
    H(doc, "5.6  Future work", 2)
    P(doc, "Future work includes adding 0.5 and 0.2 mL measurements, adding a low-cost pressure sensor to measure "
           "the p-V curve directly, using a velocity step instead of a position ramp, automating the video "
           "tracking, building a tunable-volume bubble chamber for closed-loop stiffness control, and studying "
           "temperature, bubble position, and long-term leakage.")
    PAGEBREAK(doc)

    # 6 conclusion
    H(doc, "6  Conclusion", 1)
    P(doc, "This report studied, with real measurements, how a trapped air bubble affects the static stiffness, "
           "dynamic response, and variable-stiffness capability of a syringe hydraulic transmission.")
    B(doc, "K_air falls as 1/V: 6.82±1.60, 3.41±0.13, and 1.38±0.02 N/mm at 1, 2, and 5 mL. The 1/K_air-V "
           "relation is linear (three-point fit, R² = 1.0000), with C = 6.95 N·mm/mL, V0 = 0.03 mL, and n = 1.00.")
    B(doc, "tau rises linearly with V: 0.113±0.001, 0.181±0.013, and 0.409±0.029 s; tau = 0.035 + 0.074V "
           "(three-point fit, R² = 0.9995).")
    B(doc, "The output spring and bubble in parallel give K_eff = 7.44±0.25, 4.57±0.09, and 2.40±0.05 N/mm, a "
           "3.1x range over 1-5 mL, agreeing with K_eff = k_s + K_air within 5%.")
    B(doc, "The static, dynamic, and soft-joint results are mutually consistent; more bubble means a softer and "
           "slower system. The effect can be used for variable-stiffness soft joints and for tau-based bubble "
           "detection.")
    B(doc, "The main limitations are the 1 mL resolution, the input ramp, the bubble position, and the absence of "
           "pressure and temperature measurements; 0.5 mL and below are extrapolated.")
    P(doc, "Overall, a trapped bubble in a small hydraulic system is both a passive defect and a quantifiable "
           "compliance element. With controlled bubble volume and a fixed measurement protocol, the bubble "
           "becomes a simple, low-cost, valve-free variable-stiffness knob.")
    PAGEBREAK(doc)

    # references
    H(doc, "References", 1)
    refs = [
        "Rus D, Tolley M T. Design, fabrication and control of soft robots. Nature, 2015, 521(7553): 467-475.",
        "Polygerinos P, Correll N, Morin S A, et al. Soft robotics: Review of fluid-driven intrinsically soft "
        "devices; manufacturing, sensing, control, and applications in human-robot interaction. Advanced "
        "Engineering Materials, 2017, 19(12): 1700016.",
        "Van Ham R, Sugar T G, Vanderborght B, et al. Compliant actuator designs. IEEE Robotics & Automation "
        "Magazine, 2009, 16(3): 81-94.",
        "Manring N D, Fales R C. Hydraulic Control Systems. 2nd ed. Hoboken: Wiley, 2019.",
        "Wehner M, Truby R L, Fitzgerald D J, et al. An integrated design and fabrication strategy for entirely "
        "soft, autonomous robots. Nature, 2016, 536(7617): 451-455.",
        "Manti M, Cacucciolo V, Cianchetti M. Stiffening in soft robotics: A review of the state of the art. "
        "IEEE Robotics & Automation Magazine, 2016, 23(3): 93-106.",
    ]
    for i, s in enumerate(refs, 1):
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.25
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.first_line_indent = Cm(-0.8)
        r = p.add_run(f"[{i}] {s}")
        set_font(r, "Times New Roman", 9.5)

    H(doc, "Data and code availability", 1)
    P(doc, "The measured CSV files, analysis scripts, and figure-generation scripts are available in the "
           "accompanying GitHub repository. The repository contains the raw step-response files, the static and "
           "soft-joint data, and the scripts used to generate every figure and table in this report. A DOI will "
           "be assigned through Zenodo when the repository is released.")
    P(doc, "The author declares no conflict of interest. No human or animal subjects were involved.")

    subscriptify(doc)
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    doc.save(OUT)
    print("generated ->", OUT)


if __name__ == "__main__":
    main()
